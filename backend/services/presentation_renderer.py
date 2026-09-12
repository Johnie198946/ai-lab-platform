"""Bounded editable PPTX generation and honest PDF rendering."""

from __future__ import annotations

import json
import math
import shutil
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from backend.services.presentation_scenario import DEFAULT_THEME, validate_theme

_LAYOUT_FIELDS = {
    "title": {"layout", "title", "subtitle"},
    "section": {"layout", "title"},
    "bullets": {"layout", "title", "subtitle", "bullets"},
    "conclusion": {"layout", "title", "subtitle", "bullets"},
    "two_column": {"layout", "title", "subtitle", "left", "right"},
    "table": {"layout", "title", "headers", "rows"},
    "chart": {"layout", "title", "categories", "series"},
}


def _theme(value: dict[str, Any]) -> dict[str, Any]:
    raw = validate_theme(value.get("theme", DEFAULT_THEME))
    return {
        "colors": {
            key: RGBColor.from_string(color[1:]) for key, color in raw["colors"].items()
        },
        "fonts": raw["fonts"],
    }


def _spec(content: str) -> dict[str, Any]:
    try:
        value = json.loads(content)
    except json.JSONDecodeError as exc:
        raise ValueError("presentation artifact must be valid JSON") from exc
    slides = value.get("slides") if isinstance(value, dict) else None
    if not isinstance(slides, list) or not 1 <= len(slides) <= 60:
        raise ValueError("presentation needs between 1 and 60 slides")
    if set(value) - {"title", "subtitle", "theme", "slides"}:
        raise ValueError("presentation contains unsupported fields")
    for index, slide in enumerate(slides, 1):
        if not isinstance(slide, dict):
            raise ValueError(f"slide {index} is invalid")
        layout = str(slide.get("layout") or "bullets")
        if layout not in _LAYOUT_FIELDS:
            raise ValueError(f"slide {index} has unsupported layout {layout}")
        if set(slide) - _LAYOUT_FIELDS[layout]:
            raise ValueError(f"slide {index} contains fields unused by {layout} layout")
        for key, limit in (("title", 180), ("subtitle", 240)):
            if key in slide and len(str(slide[key])) > limit:
                raise ValueError(f"slide {index} {key} exceeds {limit} characters")
        for key, count, length in (
            ("bullets", 12, 500),
            ("left", 12, 500),
            ("right", 12, 500),
            ("headers", 8, 100),
            ("categories", 20, 80),
        ):
            if key not in slide:
                continue
            items = slide[key]
            if not isinstance(items, list) or len(items) > count:
                raise ValueError(f"slide {index} {key} exceeds {count} items")
            if any(len(str(item)) > length for item in items):
                raise ValueError(
                    f"slide {index} {key} item exceeds {length} characters"
                )
        if layout == "table":
            headers, rows = slide.get("headers"), slide.get("rows")
            if (
                not isinstance(headers, list)
                or not headers
                or not isinstance(rows, list)
                or len(rows) > 12
            ):
                raise ValueError(f"slide {index} table is invalid or exceeds 12 rows")
            if any(
                not isinstance(row, list)
                or len(row) != len(headers)
                or any(len(str(cell)) > 160 for cell in row)
                for row in rows
            ):
                raise ValueError(
                    f"slide {index} table rows must match headers and cells must not exceed 160 characters"
                )
        if layout == "chart":
            categories, series = slide.get("categories"), slide.get("series")
            if (
                not isinstance(categories, list)
                or not categories
                or not isinstance(series, list)
                or not 1 <= len(series) <= 6
            ):
                raise ValueError(f"slide {index} chart is invalid or exceeds 6 series")
            for item in series:
                values = item.get("values") if isinstance(item, dict) else None
                if (
                    not isinstance(values, list)
                    or len(values) != len(categories)
                    or len(str(item.get("name") or "系列")) > 80
                ):
                    raise ValueError(
                        f"slide {index} chart series must match categories"
                    )
                try:
                    if any(not math.isfinite(float(number)) for number in values):
                        raise ValueError
                except (TypeError, ValueError) as exc:
                    raise ValueError(
                        f"slide {index} chart values must be finite numbers"
                    ) from exc
    return value


def _text(
    shape,
    theme: dict[str, Any],
    size: int = 22,
    color: str = "text",
    bold: bool = False,
    font: str = "body",
) -> None:
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = theme["fonts"][font]
            run.font.size = Pt(size)
            run.font.color.rgb = theme["colors"][color]
            run.font.bold = bold


def _title(slide, title: str, theme: dict[str, Any], subtitle: str = "") -> None:
    box = slide.shapes.add_textbox(
        Inches(0.65), Inches(0.42), Inches(11.9), Inches(0.8)
    )
    box.text_frame.text = title
    _text(box, theme, 28, bold=True, font="title")
    line = slide.shapes.add_shape(
        1, Inches(0.65), Inches(1.22), Inches(1.15), Inches(0.08)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = theme["colors"]["primary"]
    line.line.fill.background()
    if subtitle:
        sub = slide.shapes.add_textbox(
            Inches(0.67), Inches(1.4), Inches(11.4), Inches(0.55)
        )
        sub.text_frame.text = subtitle
        _text(sub, theme, 14, "muted")


def _bullets(
    slide,
    items: list[Any],
    theme: dict[str, Any],
    *,
    x: float = 0.8,
    y: float = 1.8,
    w: float = 11.6,
    h: float = 4.8,
) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = str(item)
        paragraph.level = 0
        paragraph.space_after = Pt(12)
        paragraph.font.name = theme["fonts"]["body"]
        paragraph.font.size = Pt(20)
        paragraph.font.color.rgb = theme["colors"]["text"]
    _text(box, theme, 20)


def _table(slide, spec: dict[str, Any], theme: dict[str, Any]) -> None:
    rows = spec.get("rows") or []
    headers = spec.get("headers") or []
    if not isinstance(headers, list) or not headers or not isinstance(rows, list):
        _bullets(slide, spec.get("bullets") or [], theme)
        return
    cols = len(headers)
    body = rows
    table = slide.shapes.add_table(
        len(body) + 1, cols, Inches(0.7), Inches(1.7), Inches(11.9), Inches(4.9)
    ).table
    for col, value in enumerate(headers):
        table.cell(0, col).text = str(value)
    for r, row in enumerate(body, 1):
        for c, value in enumerate(row):
            table.cell(r, c).text = str(value)
    for r in range(len(body) + 1):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = theme["colors"]["primary" if r == 0 else "pale"]
            for p in cell.text_frame.paragraphs:
                p.font.name = theme["fonts"]["body"]
                p.font.size = Pt(13)
                p.font.color.rgb = theme["colors"]["inverse" if r == 0 else "text"]
                for run in p.runs:
                    run.font.name = theme["fonts"]["body"]
                    run.font.size = Pt(13)
                    run.font.color.rgb = theme["colors"][
                        "inverse" if r == 0 else "text"
                    ]


def _chart(slide, spec: dict[str, Any], theme: dict[str, Any]) -> None:
    categories = spec.get("categories") or []
    series = spec.get("series") or []
    if (
        not isinstance(categories, list)
        or not categories
        or not isinstance(series, list)
    ):
        _bullets(slide, spec.get("bullets") or [], theme)
        return
    data = ChartData()
    data.categories = [str(x) for x in categories]
    series_count = 0
    for item in series:
        if isinstance(item, dict) and isinstance(item.get("values"), list):
            values = [float(x) for x in item["values"]]
            if len(values) == len(data.categories):
                data.add_series(str(item.get("name") or "系列"), values)
                series_count += 1
    if not series_count:
        _bullets(slide, spec.get("bullets") or [], theme)
        return
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.8),
        Inches(1.7),
        Inches(11.5),
        Inches(4.9),
        data,
    ).chart
    chart.has_legend = len(series) > 1
    chart.value_axis.has_major_gridlines = True
    for item in chart.series:
        item.format.fill.solid()
        item.format.fill.fore_color.rgb = theme["colors"]["primary"]
    for axis in (chart.category_axis, chart.value_axis):
        axis.tick_labels.font.name = theme["fonts"]["body"]
        axis.tick_labels.font.color.rgb = theme["colors"]["text"]
    if chart.has_legend:
        chart.legend.font.name = theme["fonts"]["body"]


def build_pptx(content: str) -> bytes:
    value = _spec(content)
    theme = _theme(value)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for index, item in enumerate(value["slides"]):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = theme["colors"]["background"]
        kind = str(item.get("layout") or "bullets")
        title = str(item.get("title") or f"第 {index + 1} 页")
        if kind == "title":
            box = slide.shapes.add_textbox(
                Inches(1), Inches(2.1), Inches(11.3), Inches(1.3)
            )
            box.text_frame.text = title
            _text(box, theme, 36, bold=True, font="title")
            box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            sub = slide.shapes.add_textbox(
                Inches(1.3), Inches(3.65), Inches(10.7), Inches(0.8)
            )
            sub.text_frame.text = str(
                item.get("subtitle") or value.get("subtitle") or ""
            )
            _text(sub, theme, 18, "muted")
            sub.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        elif kind == "section":
            background.fore_color.rgb = theme["colors"]["primary"]
            box = slide.shapes.add_textbox(
                Inches(1), Inches(2.4), Inches(11.3), Inches(1.4)
            )
            box.text_frame.text = title
            _text(box, theme, 36, "inverse", True, "title")
        elif kind == "two_column":
            _title(slide, title, theme, str(item.get("subtitle") or ""))
            _bullets(slide, item.get("left") or [], theme, x=0.7, w=5.7)
            _bullets(slide, item.get("right") or [], theme, x=6.9, w=5.7)
        elif kind == "chart":
            _title(slide, title, theme)
            _chart(slide, item, theme)
        elif kind == "table":
            _title(slide, title, theme)
            _table(slide, item, theme)
        else:
            _title(slide, title, theme, str(item.get("subtitle") or ""))
            _bullets(slide, item.get("bullets") or [], theme)
        number = slide.shapes.add_textbox(
            Inches(12.25), Inches(7.02), Inches(0.45), Inches(0.25)
        )
        number.text_frame.text = str(index + 1)
        _text(number, theme, 9, "muted")
    buffer = BytesIO()
    prs.save(buffer)
    data = buffer.getvalue()
    Presentation(BytesIO(data))
    return data


def render_pptx_pdf(pptx_path: Path) -> bytes:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        raise RuntimeError("PPTX 预览渲染器不可用（未检测到 LibreOffice）")
    with tempfile.TemporaryDirectory(prefix="ppt-render-") as output:
        profile = (Path(output) / "profile").as_uri()
        result = subprocess.run(
            [
                executable,
                f"-env:UserInstallation={profile}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                output,
                str(pptx_path),
            ],
            capture_output=True,
            text=True,
            timeout=90,
        )
        path = Path(output) / f"{pptx_path.stem}.pdf"
        if result.returncode or not path.is_file():
            raise RuntimeError(
                f"PPTX 预览渲染失败（exit {result.returncode}）：{(result.stderr or result.stdout).strip()[:240]}"
            )
        data = path.read_bytes()
        if not data.startswith(b"%PDF-"):
            raise RuntimeError("PPTX 预览渲染器未生成有效 PDF")
        return data
