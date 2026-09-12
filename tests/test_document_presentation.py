import hashlib
import asyncio
import json
from datetime import datetime, timedelta, timezone
from io import BytesIO
from uuid import uuid4

import pytest
import httpx
from fastapi import FastAPI
from docx import Document
from pptx import Presentation
from pypdf import PdfWriter

from backend.services import document_sources
from backend.services.document_sources import DocumentSourceError
from backend.services.presentation_renderer import build_pptx, render_pptx_pdf
from backend.services.presentation_scenario import build_presentation_plan
from backend.services.dsl_safety_compiler import DSLSafetyCompiler
from backend.services.workflow_executor import trusted_task_agent_config
from backend.api.auth import require_auth
from backend.api.documents import router as documents_router


THEME_A = {
    "colors": {
        "primary": "#1A73E8",
        "text": "#102030",
        "muted": "#405060",
        "pale": "#DDEEFF",
        "background": "#FAFBFC",
        "inverse": "#FFFFFF",
    },
    "fonts": {"title": "Arial", "body": "Courier New"},
}
THEME_B = {
    "colors": {
        "primary": "#D93025",
        "text": "#302010",
        "muted": "#605040",
        "pale": "#FFEEDD",
        "background": "#FFFDFC",
        "inverse": "#FFFF00",
    },
    "fonts": {"title": "Georgia", "body": "Times New Roman"},
}


def _docx(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _request(app: FastAPI, method: str, path: str, **kwargs):
    async def run():
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://testserver"
        ) as client:
            return await client.request(method, path, **kwargs)

    return asyncio.run(run())


def test_private_document_is_tenant_bound_and_original_survives_parse_failure(
    tmp_path, monkeypatch
):
    monkeypatch.setattr(
        document_sources,
        "note_directory",
        lambda tenant, user: tmp_path / tenant / user,
    )
    data = _docx("季度收入增长，客户留存改善。")
    receipt = document_sources.save_document_source(
        tenant_key="tenant-a",
        user_id="user-a",
        filename="report.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        data=data,
        expected_hash=hashlib.sha256(data).hexdigest(),
    )
    assert receipt["status"] == "ready"
    assert document_sources.document_text("tenant-a", "user-a", receipt["source_id"])[
        0
    ].startswith("季度收入")
    extracted = (
        document_sources.note_directory("tenant-a", "user-a")
        / ".documents"
        / receipt["source_id"]
        / "extracted.txt"
    )
    extracted.write_text("被篡改的提取内容", encoding="utf-8")
    with pytest.raises(DocumentSourceError, match="完整性") as integrity:
        document_sources.document_text("tenant-a", "user-a", receipt["source_id"])
    assert integrity.value.code == "document_integrity_error"
    with pytest.raises(DocumentSourceError, match="文档不存在"):
        document_sources.read_document_receipt(
            "tenant-a", "user-b", receipt["source_id"]
        )

    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    output = BytesIO()
    writer.write(output)
    failed = document_sources.save_document_source(
        tenant_key="tenant-a",
        user_id="user-a",
        filename="scan.pdf",
        content_type="application/pdf",
        data=output.getvalue(),
    )
    assert failed["status"] == "parse_failed"
    original, _ = document_sources.document_original_path(
        "tenant-a", "user-a", failed["source_id"]
    )
    assert original.read_bytes() == output.getvalue()
    assert failed["parse_error"]["code"] == "no_extractable_text"
    encrypted_writer = PdfWriter()
    encrypted_writer.add_blank_page(width=100, height=100)
    encrypted_writer.encrypt("secret")
    encrypted_output = BytesIO()
    encrypted_writer.write(encrypted_output)
    encrypted = document_sources.save_document_source(
        tenant_key="tenant-a",
        user_id="user-a",
        filename="locked.pdf",
        content_type="application/pdf",
        data=encrypted_output.getvalue(),
    )
    assert (
        encrypted["status"] == "parse_failed"
        and encrypted["parse_error"]["code"] == "encrypted_pdf"
    )


def test_task_agent_manifest_is_projected_to_strict_bridge_schema():
    agent = type("Agent", (), {
        "id": "agent-private",
        "private_prompt_delta": "approved prompt",
        "composition_manifest": {
            "capability_agent_ids": ["main_agent"],
            "invoked_agent_ids": ["tenant_specialist"],
            "delegation": {"max_concurrent_children": 3, "max_spawn_depth": 1},
            "knowledge_scope": ["knowledge/product/public"],
            "plan_id": "must-not-cross-runtime-boundary",
        },
    })()
    assert trusted_task_agent_config(agent) == {  # type: ignore[arg-type]
        "id": "agent-private",
        "prompt": "approved prompt",
        "capability_agent_ids": ["main_agent", "tenant_specialist"],
        "knowledge_scope": ["knowledge/product/public"],
        "delegation": {"max_concurrent_children": 3, "max_spawn_depth": 1},
    }


def test_document_rejects_legacy_and_oversize(tmp_path, monkeypatch):
    monkeypatch.setattr(
        document_sources,
        "note_directory",
        lambda tenant, user: tmp_path / tenant / user,
    )
    with pytest.raises(DocumentSourceError) as legacy:
        document_sources.save_document_source(
            tenant_key="t",
            user_id="u",
            filename="old.doc",
            content_type="application/msword",
            data=b"x",
        )
    assert legacy.value.code == "legacy_doc_unsupported"
    with pytest.raises(DocumentSourceError) as oversized:
        document_sources.save_document_source(
            tenant_key="t",
            user_id="u",
            filename="large.pdf",
            content_type="application/pdf",
            data=b"x" * (document_sources.MAX_DOCUMENT_BYTES + 1),
        )
    assert oversized.value.code == "document_too_large"


def test_authenticated_upload_receipt_download_hash_and_user_boundary(
    tmp_path, monkeypatch
):
    monkeypatch.setattr(
        document_sources,
        "note_directory",
        lambda tenant, user: tmp_path / tenant / user,
    )
    app = FastAPI()
    app.include_router(documents_router)
    identity = {"tenant_key": "tenant-a", "user_id": "user-a"}
    app.dependency_overrides[require_auth] = lambda: identity
    data = _docx("可验证的私有上传")
    digest = hashlib.sha256(data).hexdigest()
    response = _request(
        app,
        "POST",
        "/api/v1/documents",
        content=data,
        headers={
            "X-File-Name": "private.docx",
            "X-Content-Hash": digest,
            "X-File-Opt-Out": "true",
            "Content-Type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        },
    )
    assert response.status_code == 201
    receipt = response.json()
    assert (
        receipt["content_hash"] == digest
        and receipt["contribution_status"] == "excluded"
    )
    downloaded = _request(
        app, "GET", f"/api/v1/documents/{receipt['source_id']}/download"
    )
    assert (
        downloaded.content == data and downloaded.headers["x-content-sha256"] == digest
    )
    identity["user_id"] = "user-b"
    assert (
        _request(app, "GET", f"/api/v1/documents/{receipt['source_id']}").status_code
        == 404
    )


def test_authorized_default_upload_uses_authenticated_identity_and_real_queue_receipt(
    tmp_path, monkeypatch
):
    from agreement_fixtures import set_user_contribution_consent
    from backend.db import SessionLocal
    from backend.models.knowledge_contribution import KnowledgeContributionOutbox
    from backend.services.knowledge_contribution import set_contribution_policy

    monkeypatch.setattr(
        document_sources,
        "note_directory",
        lambda tenant, user: tmp_path / tenant / user,
    )
    monkeypatch.setenv("HERMES_CHAT_RUN_DB", str(tmp_path / "runs.sqlite3"))
    tenant, user = "document-" + uuid4().hex, "user-" + uuid4().hex

    async def authorize():
        await set_contribution_policy(
            tenant_key=tenant,
            enabled=True,
            agreement_version="v4",
            effective_at=datetime.now(timezone.utc) - timedelta(minutes=1),
        )
        await set_user_contribution_consent(
            tenant_key=tenant,
            user_id=user,
            service_agreement_version="service-2026-09-06",
            participation_enabled=True,
        )

    asyncio.run(authorize())
    app = FastAPI()
    app.include_router(documents_router)
    app.dependency_overrides[require_auth] = lambda: {"tenant_key": tenant, "sub": user}
    response = _request(
        app,
        "POST",
        "/api/v1/documents",
        content=_docx("参与贡献但原件保持私有"),
        headers={
            "X-File-Name": "governed.docx",
            "Content-Type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        },
    )
    receipt = response.json()
    assert response.status_code == 201 and receipt["contribution_status"] == "queued"

    async def load_event():
        async with SessionLocal() as db:
            return await db.get(
                KnowledgeContributionOutbox, receipt["contribution_event_id"]
            )

    event = asyncio.run(load_event())
    assert event and (event.tenant_key, event.user_id, event.source_id) == (
        tenant,
        user,
        receipt["source_id"],
    )
    assert receipt["contribution_run_id"]


def test_default_upload_without_server_authorization_is_denied_but_remains_private(
    tmp_path, monkeypatch
):
    monkeypatch.setattr(
        document_sources,
        "note_directory",
        lambda tenant, user: tmp_path / tenant / user,
    )
    app = FastAPI()
    app.include_router(documents_router)
    identity = {
        "tenant_key": "unauthorized-" + uuid4().hex,
        "user_id": "user-" + uuid4().hex,
    }
    app.dependency_overrides[require_auth] = lambda: identity
    receipt = _request(
        app,
        "POST",
        "/api/v1/documents",
        content=_docx("未授权不会贡献"),
        headers={
            "X-File-Name": "denied.docx",
            "Content-Type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        },
    ).json()
    assert receipt["contribution_status"] == "denied"
    assert (
        _request(
            app, "GET", f"/api/v1/documents/{receipt['source_id']}/download"
        ).status_code
        == 200
    )


@pytest.mark.parametrize(
    "result,status",
    [
        (None, "denied"),
        ({}, "failed"),
        (
            {
                "event_id": "contrib-pending",
                "status": "pending",
                "schedule_status": "pending",
                "schedule_error": "queue unavailable",
            },
            "pending",
        ),
    ],
)
def test_upload_never_reports_queued_without_accepted_scheduled_receipt(
    tmp_path, monkeypatch, result, status
):
    import backend.api.documents as api

    monkeypatch.setattr(
        document_sources,
        "note_directory",
        lambda tenant, user: tmp_path / tenant / user,
    )

    async def rejected(candidate, *, source_content):
        return result

    monkeypatch.setattr(api, "enqueue_and_schedule", rejected)
    app = FastAPI()
    app.include_router(documents_router)
    app.dependency_overrides[require_auth] = lambda: {
        "tenant_key": "tenant-no-auth",
        "user_id": "user-no-auth",
    }
    receipt = _request(
        app,
        "POST",
        "/api/v1/documents",
        content=_docx("私有保存独立成功"),
        headers={
            "X-File-Name": "private.docx",
            "Content-Type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        },
    ).json()
    assert (
        _request(
            app, "GET", f"/api/v1/documents/{receipt['source_id']}/download"
        ).status_code
        == 200
    )
    assert receipt["status"] == "ready" and receipt["contribution_status"] == status


def test_upload_contribution_failure_does_not_fail_private_save(tmp_path, monkeypatch):
    import backend.api.documents as api

    monkeypatch.setattr(
        document_sources,
        "note_directory",
        lambda tenant, user: tmp_path / tenant / user,
    )

    async def failed(candidate, *, source_content):
        raise RuntimeError("enqueue failed")

    monkeypatch.setattr(api, "enqueue_and_schedule", failed)
    app = FastAPI()
    app.include_router(documents_router)
    app.dependency_overrides[require_auth] = lambda: {
        "tenant_key": "tenant-fail",
        "user_id": "user-fail",
    }
    receipt = _request(
        app,
        "POST",
        "/api/v1/documents",
        content=_docx("仍保存"),
        headers={
            "X-File-Name": "saved.docx",
            "Content-Type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        },
    ).json()
    assert (
        _request(
            app, "GET", f"/api/v1/documents/{receipt['source_id']}/download"
        ).status_code
        == 200
    )
    assert (
        receipt["status"] == "ready"
        and receipt["contribution_status"] == "failed"
        and "enqueue failed" in receipt["contribution_error"]
    )


def test_editable_pptx_reopens_and_pdf_preview_uses_same_file(tmp_path):
    spec = {
        "title": "经营复盘",
        "slides": [
            {"layout": "title", "title": "经营复盘", "subtitle": "2026 Q3"},
            {
                "layout": "bullets",
                "title": "关键结论",
                "bullets": ["收入增长", "留存改善"],
            },
            {
                "layout": "two_column",
                "title": "对照",
                "left": ["现状"],
                "right": ["下一步"],
            },
            {
                "layout": "table",
                "title": "指标",
                "headers": ["指标", "值"],
                "rows": [["收入", "100"]],
            },
            {
                "layout": "chart",
                "title": "趋势",
                "categories": ["Q1", "Q2"],
                "series": [{"name": "收入", "values": [80, 100]}],
            },
            {"layout": "conclusion", "title": "结论", "bullets": ["聚焦增长"]},
        ],
    }
    data = build_pptx(json.dumps(spec, ensure_ascii=False))
    assert data.startswith(b"PK")
    reopened = Presentation(BytesIO(data))
    assert len(reopened.slides) == 6
    path = tmp_path / "deck.pptx"
    path.write_bytes(data)
    try:
        pdf = render_pptx_pdf(path)
    except RuntimeError as exc:
        if "预览渲染" in str(exc):
            pytest.skip(str(exc))
        raise
    assert pdf.startswith(b"%PDF-")


def test_two_real_pptx_files_apply_distinct_theme_to_title_body_table_chart_and_section(
    tmp_path,
):
    slides = [
        {"layout": "title", "title": "主题标题", "subtitle": "副标题"},
        {"layout": "bullets", "title": "正文页", "bullets": ["正文内容"]},
        {"layout": "table", "title": "表格页", "headers": ["指标"], "rows": [["收入"]]},
        {
            "layout": "chart",
            "title": "图表页",
            "categories": ["Q1"],
            "series": [{"name": "收入", "values": [1]}],
        },
        {"layout": "section", "title": "章节页"},
    ]
    decks = []
    for name, theme in (("blue", THEME_A), ("red", THEME_B)):
        path = tmp_path / f"{name}.pptx"
        path.write_bytes(
            build_pptx(
                json.dumps(
                    {"title": name, "theme": theme, "slides": slides},
                    ensure_ascii=False,
                )
            )
        )
        decks.append(Presentation(path))
    for deck, theme in zip(decks, (THEME_A, THEME_B)):
        title_run = deck.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
        body_run = (
            next(
                shape
                for shape in deck.slides[1].shapes
                if getattr(shape, "has_text_frame", False) and "正文内容" in shape.text
            )
            .text_frame.paragraphs[0]
            .runs[0]
        )
        table = next(shape.table for shape in deck.slides[2].shapes if shape.has_table)
        chart = next(shape.chart for shape in deck.slides[3].shapes if shape.has_chart)
        section_run = deck.slides[4].shapes[0].text_frame.paragraphs[0].runs[0]
        assert (
            title_run.font.name == theme["fonts"]["title"]
            and str(title_run.font.color.rgb) == theme["colors"]["text"][1:]
        )
        assert (
            body_run.font.name == theme["fonts"]["body"]
            and str(body_run.font.color.rgb) == theme["colors"]["text"][1:]
        )
        assert (
            str(table.cell(0, 0).fill.fore_color.rgb) == theme["colors"]["primary"][1:]
        )
        assert (
            table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.name
            == theme["fonts"]["body"]
        )
        assert (
            str(chart.series[0].format.fill.fore_color.rgb)
            == theme["colors"]["primary"][1:]
        )
        assert (
            str(deck.slides[4].background.fill.fore_color.rgb)
            == theme["colors"]["primary"][1:]
        )
        assert (
            section_run.font.name == theme["fonts"]["title"]
            and str(section_run.font.color.rgb) == theme["colors"]["inverse"][1:]
        )
    assert (
        decks[0].slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.name
        != decks[1].slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.name
    )


def test_theme_schema_rejects_unknown_or_incomplete_fields():
    with pytest.raises(ValueError, match="unsupported"):
        build_pptx(
            json.dumps(
                {
                    "theme": {**THEME_A, "style": "ignored"},
                    "slides": [{"layout": "title", "title": "x"}],
                }
            )
        )
    incomplete = {"colors": {"primary": "#000000"}, "fonts": THEME_A["fonts"]}
    with pytest.raises(ValueError, match="incomplete"):
        build_pptx(
            json.dumps(
                {"theme": incomplete, "slides": [{"layout": "title", "title": "x"}]}
            )
        )


def test_bridge_theme_validation_stays_inside_bridge_worker_dependencies():
    from pathlib import Path

    bridge_source = Path("scripts/hermes_bridge.py").read_text()
    scenario_source = Path("backend/services/presentation_scenario.py").read_text()
    assert "presentation_scenario import validate_theme" in bridge_source
    assert "from pptx" not in scenario_source


@pytest.mark.parametrize(
    "slide",
    [
        {"layout": "bullets", "title": "x", "bullets": ["item"] * 13},
        {"layout": "bullets", "title": "x" * 181, "bullets": []},
        {
            "layout": "table",
            "title": "x",
            "headers": ["a", "b"],
            "rows": [["only one"]],
        },
        {"layout": "unknown", "title": "x"},
    ],
)
def test_presentation_renderer_rejects_content_it_cannot_render_without_truncation(
    slide,
):
    with pytest.raises(ValueError):
        build_pptx(json.dumps({"slides": [slide]}))


def test_presentation_prompt_rejects_oversized_private_source_instead_of_excerpting():
    import scripts.hermes_bridge as bridge

    node = {
        "id": "presentation_outline",
        "node_type": "LLM_INFERENCE",
        "name": "outline",
        "parameters": {"output_format": "presentation_outline", "max_tokens": 8000},
    }
    run = {
        "goal": "deck",
        "deliverable": "pptx",
        "plan": {"nodes": [node], "edges": []},
        "nodes": {},
        "source_document": {"filename": "private.docx", "text": "x" * 8001},
    }
    with pytest.raises(RuntimeError, match="禁止静默截断"):
        bridge._workflow_node_prompt(run, node)


def test_bridge_normalizes_outline_style_points_for_renderable_slides():
    import scripts.hermes_bridge as bridge

    normalized = json.loads(bridge._normalize_presentation_reply(json.dumps({
        "title": "design",
        "theme": THEME_A,
        "slides": [
            {"layout": "bullets", "title": "goals", "key_points": ["one"]},
            {"layout": "section", "title": "flow", "subtitle": "one → two"},
            {"layout": "section", "title": "checkpoints", "key_points": ["upload", "preview"]},
            {"layout": "process", "title": "steps", "steps": ["one", "two"]},
        ],
    })))
    assert normalized["slides"][0]["bullets"] == ["one"]
    assert "key_points" not in normalized["slides"][0]
    assert normalized["slides"][1]["layout"] == "title"
    assert normalized["slides"][2]["layout"] == "bullets"
    assert normalized["slides"][2]["bullets"] == ["upload", "preview"]
    assert normalized["slides"][3]["layout"] == "bullets"
    assert normalized["slides"][3]["bullets"] == ["one", "two"]
    assert "steps" not in normalized["slides"][3]
    build_pptx(json.dumps(normalized))


def test_bridge_final_projection_uses_exact_approved_design_theme_and_binding():
    import scripts.hermes_bridge as bridge

    design = json.dumps(
        {
            "title": "approved",
            "theme": THEME_A,
            "slides": [{"layout": "title", "title": "sample"}],
        },
        ensure_ascii=False,
        separators=(",", ":"),
    )
    digest = hashlib.sha256(design.encode()).hexdigest()
    run = {
        "plan": {
            "nodes": [
                {
                    "id": "presentation_design",
                    "parameters": {"output_format": "presentation_design"},
                }
            ]
        },
        "nodes": {
            "presentation_design": {
                "status": "succeeded",
                "attempt": 3,
                "output": design,
            }
        },
        "approved_gates": ["presentation_design"],
        "approved_gate_artifacts": {
            "presentation_design": {
                "artifact_id": "wfa_" + "a" * 32,
                "content_hash": digest,
                "artifact_version": 3,
            }
        },
    }
    model_final = json.dumps(
        {
            "theme": THEME_B,
            "slides": [{"layout": "bullets", "title": "final", "bullets": ["kept"]}],
        }
    )
    projected, binding = bridge._bind_approved_presentation_design(run, model_final)
    assert json.loads(projected)["theme"] == {
        "colors": {key: value.upper() for key, value in THEME_A["colors"].items()},
        "fonts": THEME_A["fonts"],
    }
    assert (
        json.loads(projected)["slides"][0]["bullets"] == ["kept"]
        and binding["content_hash"] == digest
    )
    run["nodes"]["presentation_design"]["output"] = design + " "
    with pytest.raises(RuntimeError, match="tampered"):
        bridge._bind_approved_presentation_design(run, model_final)


@pytest.mark.asyncio
async def test_workflow_projection_rechecks_approved_design_against_final_theme(
    tmp_path, monkeypatch
):
    from types import SimpleNamespace
    from backend.services.workflow_executor import _assert_approved_presentation_projection
    import backend.services.workflow_artifacts as artifacts

    approved = json.dumps(
        {"theme": THEME_A, "slides": [{"layout": "title", "title": "sample"}]}
    )
    path = tmp_path / "design.json"
    path.write_text(approved)
    digest = hashlib.sha256(approved.encode()).hexdigest()
    design = SimpleNamespace(
        id="wfa_" + "a" * 32,
        execution_id="exec-theme",
        content_hash=digest,
        relative_path="design.json",
    )
    outline_value = {
        "title": "final",
        "slides": [{"layout": "title", "title": "final"}],
    }
    outline_text = json.dumps(outline_value)
    outline_path = tmp_path / "outline.json"
    outline_path.write_text(outline_text)
    outline_digest = hashlib.sha256(outline_text.encode()).hexdigest()
    outline = SimpleNamespace(
        id="wfa_" + "b" * 32,
        execution_id="exec-theme",
        content_hash=outline_digest,
        relative_path="outline.json",
    )

    class DB:
        async def scalar(self, query):
            return object()

        async def get(self, model, key):
            return {design.id: design, outline.id: outline}.get(key)

    monkeypatch.setattr(artifacts, "run_root", lambda execution: tmp_path)
    binding = {"artifact_id": design.id, "content_hash": digest, "artifact_version": 2}
    outline_binding = {
        "artifact_id": outline.id,
        "content_hash": outline_digest,
        "artifact_version": 1,
    }
    execution = SimpleNamespace(id="exec-theme")
    await _assert_approved_presentation_projection(
        DB(),
        execution,
        {
            "approved_design": binding,
            "approved_outline": outline_binding,
            "content": json.dumps(
                {"theme": THEME_A, "slides": [{"layout": "title", "title": "final"}]}
            ),
        },
        "presentation",
    )
    with pytest.raises(ValueError, match="differs"):
        await _assert_approved_presentation_projection(
            DB(),
            execution,
            {
                "approved_design": binding,
                "approved_outline": outline_binding,
                "content": json.dumps(
                    {
                        "theme": THEME_B,
                        "slides": [{"layout": "title", "title": "tampered"}],
                    }
                ),
            },
            "presentation",
        )


def test_final_presentation_requires_approved_outline_binding_and_structure():
    import scripts.hermes_bridge as bridge

    design = json.dumps({"theme": THEME_A, "slides": [{"layout": "title", "title": "sample"}]})
    outline = json.dumps(
        {
            "title": "deck",
            "slides": [
                {"layout": "title", "title": "deck"},
                {"layout": "bullets", "title": "findings", "key_points": ["one"]},
            ],
        }
    )
    run = {
        "plan": {
            "nodes": [
                {"id": "presentation_outline", "parameters": {"output_format": "presentation_outline"}},
                {"id": "presentation_design", "parameters": {"output_format": "presentation_design"}},
            ]
        },
        "nodes": {
            "presentation_outline": {"attempt": 1, "output": outline},
            "presentation_design": {"attempt": 2, "output": design},
        },
        "approved_gates": ["presentation_outline", "presentation_design"],
        "approved_gate_artifacts": {
            "presentation_outline": {
                "artifact_id": "wfa_" + "a" * 32,
                "content_hash": hashlib.sha256(outline.encode()).hexdigest(),
                "artifact_version": 1,
            },
            "presentation_design": {
                "artifact_id": "wfa_" + "b" * 32,
                "content_hash": hashlib.sha256(design.encode()).hexdigest(),
                "artifact_version": 2,
            },
        },
    }
    final = json.dumps(
        {
            "title": "deck",
            "slides": [
                {"layout": "title", "title": "deck"},
                {"layout": "bullets", "title": "findings", "bullets": ["one"]},
            ],
        }
    )
    projected, design_binding, outline_binding = bridge._bind_approved_presentation_inputs(run, final)
    assert json.loads(projected)["theme"]["fonts"] == THEME_A["fonts"]
    assert design_binding["artifact_version"] == 2
    assert outline_binding["artifact_version"] == 1
    changed = json.loads(final)
    changed["slides"][1]["title"] = "unapproved"
    with pytest.raises(RuntimeError, match="approved outline structure"):
        bridge._bind_approved_presentation_inputs(run, json.dumps(changed))
    run["nodes"]["presentation_outline"]["output"] = outline + " "
    with pytest.raises(RuntimeError, match="outline.*tampered"):
        bridge._bind_approved_presentation_inputs(run, final)


def test_hermes_gate_rejects_stale_version_and_records_current_approval(monkeypatch):
    import scripts.hermes_bridge as bridge
    from fastapi import HTTPException

    output = '{"title":"outline","slides":[{"layout":"title","title":"x"}]}'
    digest = hashlib.sha256(output.encode()).hexdigest()
    run = {
        "execution_id": "exec-gate",
        "status": "awaiting_approval",
        "next_seq": 1,
        "events": [],
        "nodes": {"outline": {"status": "succeeded", "attempt": 2, "output": output}},
        "approved_gates": [],
        "approved_gate_artifacts": {},
        "plan": {
            "nodes": [{"id": "outline", "parameters": {"approval_gate": "outline"}}],
            "edges": [],
        },
    }
    monkeypatch.setattr(bridge, "HERMES_BRIDGE_INTERNAL_TOKEN", "secret")
    monkeypatch.setattr(bridge, "_start_workflow_thread", lambda execution_id: None)
    monkeypatch.setattr(bridge, "_save_workflow_runs", lambda: None)
    bridge._workflow_runs["exec-gate"] = run
    try:
        with pytest.raises(HTTPException) as stale:
            asyncio.run(
                bridge.approve_workflow_gate(
                    "exec-gate",
                    bridge.WorkflowGateApprovalRequest(
                        node_id="outline",
                        artifact_version=1,
                        artifact_id="wfa_" + "a" * 32,
                        expected_hash=digest,
                    ),
                    "secret",
                )
            )
        assert stale.value.status_code == 409
        with pytest.raises(HTTPException) as tampered:
            asyncio.run(
                bridge.approve_workflow_gate(
                    "exec-gate",
                    bridge.WorkflowGateApprovalRequest(
                        node_id="outline",
                        artifact_version=2,
                        artifact_id="wfa_" + "a" * 32,
                        expected_hash="0" * 64,
                    ),
                    "secret",
                )
            )
        assert tampered.value.status_code == 409
        result = asyncio.run(
            bridge.approve_workflow_gate(
                "exec-gate",
                bridge.WorkflowGateApprovalRequest(
                    node_id="outline",
                    artifact_version=2,
                    artifact_id="wfa_" + "a" * 32,
                    expected_hash=digest,
                ),
                "secret",
            )
        )
        assert result["status"] == "queued" and run["approved_gates"] == ["outline"]
        assert run["approved_gate_artifacts"]["outline"]["content_hash"] == digest
        retried = asyncio.run(
            bridge.retry_workflow_run(
                "exec-gate",
                bridge.WorkflowRetryRequest(from_node_id="outline"),
                "secret",
            )
        )
        assert (
            retried["status"] == "queued"
            and run["approved_gates"] == []
            and run["approved_gate_artifacts"] == {}
        )
    finally:
        bridge._workflow_runs.pop("exec-gate", None)


def test_presentation_scenario_has_two_business_gates_before_binary_deck():
    workflow = type(
        "Workflow",
        (),
        {
            "title": "Deck",
            "requirements_snapshot": {
                "scenario_id": "document-to-presentation",
                "source_document": {
                    "source_id": "doc_1",
                    "content_hash": "a" * 64,
                    "source_revision": 1,
                },
            },
        },
    )()
    plan = build_presentation_plan(workflow, plan_id="plan", knowledge_scope=[])
    assert [node["parameters"].get("approval_gate") for node in plan["nodes"]] == [
        "outline",
        "design",
        None,
    ]
    assert plan["nodes"][-1]["parameters"]["output_format"] == "presentation"
    assert {tuple(edge.values()) for edge in plan["edges"]} >= {
        ("presentation_outline", "presentation_deck"),
        ("presentation_design", "presentation_deck"),
    }
    assert len(DSLSafetyCompiler.compile_and_validate(plan).nodes) == 3
